from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = r'f:\00.work\01.brycen\workspace\AI-Review-Document-System\backend\uploads\P012_Project012.pptx'

if not os.path.exists(file_path):
    print(f"File not found: {file_path}")
    exit(1)

try:
    prs = Presentation(file_path)
    slide_3 = prs.slides[2] # 0-indexed
    print(f"Inspecting Slide 3 of {file_path}")
    print(f"Number of shapes: {len(slide_3.shapes)}")
    
    for i, shape in enumerate(slide_3.shapes, 1):
        print(f"\nShape {i}: {shape.name}")
        print(f" - Type: {shape.shape_type}")
        print(f" - Has Text Frame: {shape.has_text_frame}")
        print(f" - Has Table: {shape.has_table}")
        print(f" - Has Chart: {shape.has_chart}")
        
        if shape.has_text_frame:
            print(f" - Text: {shape.text.strip()}")
        
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f" - [PICTURE] This is an image.")
        
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            print(f" - [GROUP] Contains {len(shape.shapes)} sub-shapes.")
            
        if shape.shape_type == MSO_SHAPE_TYPE.CHART:
            print(f" - [CHART] This is a chart.")

except Exception as e:
    print(f"Error: {e}")
