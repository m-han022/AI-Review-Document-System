from openpyxl import load_workbook

def extract_text_from_xlsx(file_path: str) -> str:
    max_sheets = 10
    max_rows_per_sheet = 2000
    max_cells = 50000
    wb = load_workbook(filename=file_path, data_only=True, read_only=True)
    try:
        parts: list[str] = []
        cell_count = 0
        for sheet in wb.worksheets[:max_sheets]:
            parts.append(f"[Sheet {sheet.title}]")
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                row_count += 1
                if row_count > max_rows_per_sheet:
                    parts.append("[Truncated: max rows reached]")
                    break
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    parts.append(" | ".join(cells))
                    cell_count += len(cells)
                    if cell_count >= max_cells:
                        parts.append("[Truncated: max cells reached]")
                        return "\n".join(parts)
        return "\n".join(parts)
    finally:
        wb.close()
