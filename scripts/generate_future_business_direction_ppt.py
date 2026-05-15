from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path("artifacts") / "Future_Business_Direction_AI_Review.pptx"

WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(24, 52, 88)
BLUE = RGBColor(45, 112, 196)
TEAL = RGBColor(34, 139, 131)
GREEN = RGBColor(48, 153, 102)
AMBER = RGBColor(201, 124, 45)
GRAY = RGBColor(88, 105, 124)
LIGHT = RGBColor(241, 246, 250)


def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=""):
    add_textbox(slide, 0.7, 0.45, 11.8, 0.7, title, 24, True, NAVY)
    if subtitle:
        add_textbox(slide, 0.7, 1.05, 11.6, 0.6, subtitle, 12, False, GRAY)


def add_bullets(slide, items, left=0.9, top=1.9, width=11.0, height=4.8, size=17):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = GRAY


def add_panel(slide, left, top, width, height, title, items, line_color):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = line_color
    add_textbox(slide, left + 0.2, top + 0.18, width - 0.4, 0.35, title, 16, True, line_color)
    add_bullets(slide, items, left + 0.2, top + 0.7, width - 0.4, height - 0.9, 13.5)


def add_footer(slide, text):
    add_textbox(slide, 0.7, 7.0, 12.0, 0.22, text, 10, False, GRAY)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(
        slide,
        "Dinh huong nghiep vu du kien bo sung cho he thong AI Review",
        "Deck nay tom tat cac nang luc se mo rong trong tuong lai theo huong an toan, de van hanh, va co the rollout tung pha.",
    )
    add_bullets(
        slide,
        [
            "Chuyen tu he thong cham tai lieu thanh nen tang review, governance, va knowledge sharing.",
            "Giu UX don gian cho user, day complexity ve backend, bundle governance, va admin flow.",
            "Mo rong dan theo roadmap, khong ngat quang he thong dang chay.",
        ],
        top=2.0,
        height=2.2,
    )
    add_panel(
        slide,
        0.85,
        4.55,
        3.8,
        1.6,
        "Truc hien tai",
        ["Review theo loai tai lieu", "Cham theo bundle active", "Audit va risk warning"],
        BLUE,
    )
    add_panel(
        slide,
        4.8,
        4.55,
        3.8,
        1.6,
        "Truc mo rong",
        ["Learning case", "Reference snippets", "Universal input"],
        TEAL,
    )
    add_panel(
        slide,
        8.75,
        4.55,
        3.7,
        1.6,
        "Nguyen tac",
        ["Khong pha flow cu", "Feature flag", "Rollback duoc"],
        GREEN,
    )
    add_footer(slide, "Future business direction overview")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "1. Muc tieu nghiep vu se bo sung")
    add_bullets(
        slide,
        [
            "Khong chi review de cham diem, ma ho tro quality gate, decision support, va risk awareness.",
            "Khong chi hoc tu ket qua review, ma hoc tu noi dung tai lieu viet tot de rollout ngang toan cong ty.",
            "Khong chi doc file upload, ma co the review noi dung den tu nhieu nguon du lieu khac nhau.",
            "Khong chi phuc vu tung project, ma tong hop tri thuc cho toan cong ty theo thang.",
        ],
        top=1.95,
        height=3.4,
    )
    add_footer(slide, "Business objectives beyond scoring")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "2. Bundle governance va decision support")
    add_panel(
        slide,
        0.8,
        1.95,
        4.0,
        4.9,
        "Nhung gi se co",
        [
            "Scope = document_type + muc danh gia",
            "Bundle active duy nhat theo scope",
            "Clone -> Validate -> Activate",
            "Decision model: pass / conditional_pass / fail / needs_human_review",
        ],
        BLUE,
    )
    add_panel(
        slide,
        4.95,
        1.95,
        3.85,
        4.9,
        "Gia tri",
        [
            "De audit",
            "De rollback",
            "De doi bundle ma khong sua tay logic",
            "Giu UI don gian voi user",
        ],
        TEAL,
    )
    add_panel(
        slide,
        8.95,
        1.95,
        3.55,
        4.9,
        "Loi ich van hanh",
        [
            "Quyet dinh ro rang hon",
            "Giam lech danh gia giua reviewer",
            "PMO thay duoc bundle dang ap dung",
        ],
        GREEN,
    )
    add_footer(slide, "Evaluation governance foundation")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "3. Risk warning theo boi canh du an")
    add_bullets(
        slide,
        [
            "Bo sung project context: project_description, domain, phase, milestone, constraints, dependencies, criticality.",
            "AI khong chi cham diem tai lieu, ma canh bao risk signal co lien quan den release, dependency, compliance, handoff.",
            "Moi warning phai co severity, reason, evidence, recommended_action.",
            "Neu khong du can cu thi dua ve needs_human_review thay vi ket luan manh.",
        ],
        top=2.0,
        height=3.4,
    )
    add_panel(
        slide,
        1.0,
        5.5,
        5.2,
        1.0,
        "Nguyen tac",
        ["Chi canh bao tren du lieu that, khong co AI confidence score gia."],
        AMBER,
    )
    add_panel(
        slide,
        6.5,
        5.5,
        5.2,
        1.0,
        "Gia tri",
        ["PMO thay risk som truoc khi qua release gate."],
        BLUE,
    )
    add_footer(slide, "Context-aware risk detection")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "4. Learning case tu review result")
    add_bullets(
        slide,
        [
            "Tong hop best practices, recurring issues, risk patterns, release-readiness cases.",
            "Tao learning candidates tu review result, khong publish tu dong.",
            "Human curation truoc khi dua vao Learning Hub, monthly digest, hoac searchable knowledge base.",
            "Theo doi impact sau rollout: case nao duoc tai su dung, recurring issue co giam khong.",
        ],
        top=2.0,
        height=3.2,
    )
    add_panel(
        slide,
        1.0,
        5.35,
        4.0,
        1.1,
        "Flow",
        ["Review -> Candidate -> Curate -> Approve -> Publish"],
        TEAL,
    )
    add_panel(
        slide,
        5.2,
        5.35,
        3.2,
        1.1,
        "Output",
        ["Monthly digest", "Learning case library"],
        BLUE,
    )
    add_panel(
        slide,
        8.65,
        5.35,
        3.7,
        1.1,
        "Rule",
        ["Khong auto publish", "Can sensitivity scan"],
        GREEN,
    )
    add_footer(slide, "Learning from review intelligence")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "5. Reference snippets va reusable patterns")
    add_bullets(
        slide,
        [
            "Khong chi hoc tu ket qua cham, ma hoc tu doan noi dung input viet tot.",
            "Trich doan requirement, dependency, risk, summary, acceptance criteria co gia tri tai su dung.",
            "Moi snippet phai co: why_good, reuse_guidance, source reference, sensitivity level.",
            "Ngoai text snippet, tuong lai ho tro table snippet, field pattern, conversation excerpt.",
        ],
        top=2.0,
        height=3.3,
    )
    add_panel(
        slide,
        1.0,
        5.45,
        4.1,
        1.0,
        "Thu vien tri thuc",
        ["Snippet library", "Pattern library"],
        AMBER,
    )
    add_panel(
        slide,
        5.35,
        5.45,
        3.1,
        1.0,
        "Control",
        ["Masking", "Approval", "Visibility scope"],
        TEAL,
    )
    add_panel(
        slide,
        8.7,
        5.45,
        3.5,
        1.0,
        "Gia tri",
        ["Tai lieu mau tot de rollout ngang"],
        BLUE,
    )
    add_footer(slide, "Learning from strong input content")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "6. Universal input model")
    add_bullets(
        slide,
        [
            "Tuong lai khong chi review file PDF/PPT, ma review normalized content from any source.",
            "Tach document_type khoi input_source_type.",
            "Dua moi input qua extraction / normalization / canonical content model truoc khi review.",
            "Adapter/plugin cho PDF, PPTX, DOCX, XLSX, URL, wiki, ticket, transcript, JSON, OCR.",
        ],
        top=2.0,
        height=3.3,
    )
    add_panel(
        slide,
        0.95,
        5.45,
        4.0,
        1.0,
        "Nguyen tac",
        ["Review normalized content, khong khoa vao file format"],
        GREEN,
    )
    add_panel(
        slide,
        5.15,
        5.45,
        3.2,
        1.0,
        "Metadata",
        ["Extraction status", "Confidence", "Warnings"],
        BLUE,
    )
    add_panel(
        slide,
        8.55,
        5.45,
        3.7,
        1.0,
        "Loi ich",
        ["Mo duong cho nhieu he thong du lieu khac nhau"],
        TEAL,
    )
    add_footer(slide, "From file review to content review")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "7. External sources va connector governance")
    add_bullets(
        slide,
        [
            "Bo sung kha nang doc tu SharePoint, file server, document repository noi bo, va cac source chi dinh ben ngoai.",
            "Tach che do snapshot-first va reference-only.",
            "Review chinh thuc nen uu tien snapshot-first de replay va audit duoc.",
            "Moi lan doc nguon ngoai phai luu source_locator, retrieved_at, retrieved_by, retrieval_status.",
        ],
        top=2.0,
        height=3.2,
    )
    add_panel(
        slide,
        0.95,
        5.35,
        3.8,
        1.1,
        "Security",
        ["Permission boundary", "Allowed connector", "Allowed path/site/folder"],
        AMBER,
    )
    add_panel(
        slide,
        4.95,
        5.35,
        3.8,
        1.1,
        "Audit",
        ["Replayable snapshot", "Connector logs", "Retrieval metadata"],
        BLUE,
    )
    add_panel(
        slide,
        8.95,
        5.35,
        3.2,
        1.1,
        "UX",
        ["Source indicator", "Snapshot vs reference", "Reliability hint"],
        TEAL,
    )
    add_footer(slide, "External source reading must not break auditability")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "8. Thu tu trien khai de khong ngat quang he thong")
    add_panel(slide, 0.8, 2.0, 2.4, 3.8, "P1-P2", ["Governance", "Scope", "Bundle", "Active mapping"], BLUE)
    add_panel(slide, 3.45, 2.0, 2.4, 3.8, "P3-P4", ["Runtime contract", "Audit", "UI baseline"], TEAL)
    add_panel(slide, 6.1, 2.0, 2.4, 3.8, "P5-P7", ["Risk warning", "Learning case", "Snippets"], GREEN)
    add_panel(slide, 8.75, 2.0, 2.4, 3.8, "P8", ["Universal input", "Canonical content", "Adapters"], AMBER)
    add_panel(slide, 11.4, 2.0, 1.0, 3.8, "P9", ["External", "sources"], NAVY)
    add_bullets(
        slide,
        [
            "Rule: chi expose UI moi sau khi backend contract on dinh.",
            "Moi phase phai co default path va rollback point.",
            "Upload legacy phai tiep tuc chay trong suot qua trinh mo rong.",
        ],
        top=6.0,
        height=0.9,
        size=14,
    )
    add_footer(slide, "Low-disruption implementation roadmap")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "9. Dieu kien de rollout an toan")
    add_bullets(
        slide,
        [
            "Spec, contract, UI, va docs phai thay doi cung nhau.",
            "Khong active bundle, connector, hoac input source moi neu chua co test va rollback path.",
            "Learning candidate va snippet candidate phai qua sensitivity scan va redaction gate.",
            "Review chinh thuc tu external source phai replay duoc.",
            "Metrics phai do duoc adoption, failure rate, override rate, va impact learning.",
        ],
        top=2.0,
        height=3.4,
    )
    add_panel(
        slide,
        1.0,
        5.45,
        4.0,
        1.0,
        "Minimum safe foundation",
        ["Governance + bundle + runtime contract + UI baseline"],
        GREEN,
    )
    add_panel(
        slide,
        5.3,
        5.45,
        3.3,
        1.0,
        "Do not skip",
        ["Testing", "Audit", "Rollback"],
        AMBER,
    )
    add_panel(
        slide,
        8.9,
        5.45,
        3.3,
        1.0,
        "Owner model",
        ["PMO", "Backend", "Frontend", "Design", "QA", "Security"],
        BLUE,
    )
    add_footer(slide, "Release safety conditions for future business expansion")

    return prs


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs = build()
    prs.save(str(OUT_PATH))
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    main()
