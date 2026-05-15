from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path("artifacts") / "Future_Business_Direction_AI_Review_Executive_VI.pptx"

WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(22, 46, 79)
BLUE = RGBColor(49, 110, 194)
TEAL = RGBColor(33, 139, 128)
GREEN = RGBColor(56, 161, 105)
AMBER = RGBColor(196, 122, 49)
GRAY = RGBColor(95, 109, 125)
LIGHT = RGBColor(242, 246, 250)


def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=""):
    add_textbox(slide, 0.7, 0.45, 11.9, 0.75, title, 25, True, NAVY)
    if subtitle:
        add_textbox(slide, 0.7, 1.08, 11.5, 0.55, subtitle, 12, False, GRAY)


def add_bullets(slide, items, left=0.9, top=1.9, width=11.2, height=4.8, size=17):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = GRAY


def add_panel(slide, left, top, width, height, title, items, line_color):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = line_color
    add_textbox(slide, left + 0.18, top + 0.16, width - 0.36, 0.32, title, 15, True, line_color)
    add_bullets(slide, items, left + 0.18, top + 0.62, width - 0.36, height - 0.75, 13.2)


def add_footer(slide, text):
    add_textbox(slide, 0.7, 7.0, 12.0, 0.2, text, 10, False, GRAY)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(
        slide,
        "Định hướng mở rộng nghiệp vụ cho hệ thống AI Review",
        "Bản tóm tắt dành cho quản lý, tập trung vào giá trị kinh doanh, hướng mở rộng và cách triển khai an toàn.",
    )
    add_bullets(
        slide,
        [
            "Mở rộng từ công cụ chấm tài liệu thành nền tảng hỗ trợ đánh giá, cảnh báo rủi ro và chia sẻ tri thức.",
            "Giữ trải nghiệm đơn giản cho người dùng phổ thông, đồng thời tăng khả năng quản trị và kiểm soát cho PMO/Admin.",
            "Triển khai theo từng bước nhỏ để không làm gián đoạn hệ thống đang vận hành.",
        ],
        top=2.0,
        height=2.5,
        size=16,
    )
    add_panel(slide, 0.9, 5.1, 3.7, 1.2, "Giá trị", ["Nhanh hơn", "Đồng bộ hơn", "Kiểm soát tốt hơn"], BLUE)
    add_panel(slide, 4.8, 5.1, 3.7, 1.2, "Hướng mở rộng", ["Cảnh báo rủi ro", "Case learning", "Đa nguồn dữ liệu"], TEAL)
    add_panel(slide, 8.7, 5.1, 3.7, 1.2, "Nguyên tắc", ["Không phá flow cũ", "Có rollback", "Có kiểm soát"], GREEN)
    add_footer(slide, "Executive summary")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "1. Hệ thống trong tương lai sẽ phục vụ gì")
    add_bullets(
        slide,
        [
            "Đánh giá chất lượng tài liệu nhất quán hơn theo từng loại tài liệu và mức đánh giá.",
            "Hỗ trợ quyết định duyệt, yêu cầu chỉnh sửa hoặc cần xem xét thêm trước các mốc quan trọng.",
            "Phát hiện sớm rủi ro dựa trên bối cảnh thực tế của dự án, không chỉ dựa trên nội dung tài liệu đơn lẻ.",
            "Tận dụng kết quả review và các đoạn mô tả tốt để tạo tri thức dùng chung cho toàn công ty.",
        ],
        top=2.0,
        height=3.7,
        size=16,
    )
    add_footer(slide, "Future business capability")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "2. Những năng lực sẽ được bổ sung")
    add_panel(
        slide,
        0.8,
        1.95,
        3.8,
        4.9,
        "Quản trị đánh giá",
        [
            "Bộ tiêu chuẩn áp dụng theo loại tài liệu và mức đánh giá",
            "Thay đổi có version, có kiểm tra, có phê duyệt",
            "Dễ audit và dễ quay lui",
        ],
        BLUE,
    )
    add_panel(
        slide,
        4.8,
        1.95,
        3.8,
        4.9,
        "Cảnh báo rủi ro",
        [
            "Đọc tài liệu trong bối cảnh dự án",
            "Phát hiện sớm điểm có thể ảnh hưởng release, dependency, compliance",
            "Mỗi cảnh báo có lý do và bằng chứng",
        ],
        TEAL,
    )
    add_panel(
        slide,
        8.8,
        1.95,
        3.7,
        4.9,
        "Tri thức tổ chức",
        [
            "Case learning hàng tháng",
            "Đoạn mô tả hay để tham khảo dùng lại",
            "Mẫu tốt để triển khai ngang",
        ],
        GREEN,
    )
    add_footer(slide, "Key expansion capabilities")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "3. Hướng mở rộng nguồn dữ liệu đầu vào")
    add_bullets(
        slide,
        [
            "Tương lai không chỉ đọc file upload như PDF/PPT, mà có thể đọc nội dung từ nhiều nguồn dữ liệu khác nhau.",
            "Bao gồm: tài liệu Word/Excel, nội dung web, wiki, ticket, transcript, JSON và các nguồn bên ngoài như SharePoint hoặc file server.",
            "Hệ thống sẽ chuẩn hóa nội dung về một mô hình chung trước khi đánh giá để tránh phụ thuộc vào định dạng file.",
            "Với nguồn ngoài, ưu tiên cơ chế snapshot để đảm bảo có thể truy vết và kiểm chứng lại nội dung đã review.",
        ],
        top=2.0,
        height=3.6,
        size=16,
    )
    add_panel(slide, 1.0, 5.75, 3.5, 0.9, "Nguồn dữ liệu", ["Upload, URL, wiki, ticket, SharePoint, file server"], TEAL)
    add_panel(slide, 4.9, 5.75, 3.2, 0.9, "Nguyên tắc", ["Chuẩn hóa trước khi review"], BLUE)
    add_panel(slide, 8.45, 5.75, 3.7, 0.9, "Kiểm soát", ["Snapshot, permission, audit"], GREEN)
    add_footer(slide, "Universal input direction")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "4. Lộ trình triển khai để không làm gián đoạn hệ thống")
    add_panel(slide, 0.8, 2.0, 2.3, 3.9, "P1-P2", ["Chốt baseline", "Thiết lập scope", "Quản trị bundle"], BLUE)
    add_panel(slide, 3.35, 2.0, 2.3, 3.9, "P3-P4", ["Chuẩn hóa contract", "Audit", "UI baseline"], TEAL)
    add_panel(slide, 5.9, 2.0, 2.3, 3.9, "P5-P6", ["Risk warning", "Learning case"], GREEN)
    add_panel(slide, 8.45, 2.0, 2.3, 3.9, "P7-P8", ["Snippets", "Đa nguồn input"], AMBER)
    add_panel(slide, 11.0, 2.0, 1.5, 3.9, "P9", ["Nguồn ngoài", "Connector"], NAVY)
    add_bullets(
        slide,
        [
            "Chỉ mở rộng giao diện khi lớp backend và contract đã ổn định.",
            "Mỗi giai đoạn đều phải có đường lui và cơ chế bật/tắt riêng.",
        ],
        top=6.1,
        height=0.8,
        size=14,
    )
    add_footer(slide, "Low-disruption roadmap")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "5. Điều kiện để triển khai thành công")
    add_bullets(
        slide,
        [
            "Phải có nguồn sự thật thống nhất giữa đặc tả, backend, frontend và tài liệu trình bày.",
            "Mọi thay đổi cấu hình đánh giá phải đi qua quy trình clone, kiểm tra và kích hoạt có kiểm soát.",
            "Các phần mở rộng như risk warning, learning case, snippet sharing hay external source đều phải có kiểm soát nhạy cảm và khả năng audit.",
            "PMO, Backend, Frontend, Design, QA và Security phải có owner rõ cho từng giai đoạn triển khai.",
        ],
        top=2.0,
        height=3.6,
        size=16,
    )
    add_panel(slide, 1.0, 5.8, 5.0, 0.9, "Nền tảng tối thiểu cần có trước", ["Governance + bundle + contract runtime + UI baseline"], GREEN)
    add_panel(slide, 6.35, 5.8, 5.0, 0.9, "Thông điệp chốt", ["Mở rộng được nhưng vẫn giữ hệ thống ổn định, rõ ràng và dễ vận hành"], BLUE)
    add_footer(slide, "Success conditions")

    return prs


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs = build()
    prs.save(str(OUT_PATH))
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    main()
