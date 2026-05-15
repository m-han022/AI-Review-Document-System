from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path("artifacts") / "Future_Business_Direction_AI_Review_Executive.pptx"

WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(24, 52, 88)
BLUE = RGBColor(45, 112, 196)
TEAL = RGBColor(34, 139, 131)
GREEN = RGBColor(48, 153, 102)
AMBER = RGBColor(201, 124, 45)
GRAY = RGBColor(88, 105, 124)
LIGHT = RGBColor(241, 246, 250)


def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=""):
    add_textbox(slide, 0.7, 0.45, 11.8, 0.7, title, 24, True, NAVY)
    if subtitle:
        add_textbox(slide, 0.7, 1.05, 11.6, 0.55, subtitle, 12, False, GRAY)


def add_bullets(slide, items, left=0.9, top=1.9, width=11.2, height=4.8, size=17):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = GRAY


def add_panel(slide, left, top, width, height, title, items, line_color):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = line_color
    add_textbox(slide, left + 0.18, top + 0.16, width - 0.36, 0.32, title, 15, True, line_color)
    add_bullets(slide, items, left + 0.18, top + 0.62, width - 0.36, height - 0.75, 13.5)


def add_footer(slide, text):
    add_textbox(slide, 0.7, 7.0, 12.0, 0.2, text, 10, False, GRAY)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(
        slide,
        "Dinh huong nghiep vu tuong lai cho he thong AI Review",
        "Ban executive tom tat de trinh bay nhanh ve huong mo rong va cach rollout an toan.",
    )
    add_bullets(
        slide,
        [
            "Tu he thong cham tai lieu thanh nen tang review, risk awareness, governance, va knowledge sharing.",
            "Giu UX don gian cho user, day complexity ve bundle governance, admin flow, va audit.",
            "Mo rong dan theo pha, khong ngat quang he thong dang van hanh.",
        ],
        top=2.0,
        height=2.5,
    )
    add_panel(slide, 0.9, 5.1, 3.7, 1.2, "Gia tri", ["Nhanh hon", "Dong bo hon", "Audit duoc"], BLUE)
    add_panel(slide, 4.8, 5.1, 3.7, 1.2, "Mo rong", ["Risk warning", "Learning", "Multi-source input"], TEAL)
    add_panel(slide, 8.7, 5.1, 3.7, 1.2, "Nguyen tac", ["Khong pha flow cu", "Rollback duoc", "Feature flag"], GREEN)
    add_footer(slide, "Executive overview")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "1. Mo hinh nghiep vu se duoc nang cap")
    add_panel(
        slide,
        0.8,
        1.95,
        3.9,
        4.8,
        "Nen tang danh gia",
        [
            "Scope = document_type + muc danh gia",
            "Bundle active duy nhat theo scope",
            "Decision model ro rang",
        ],
        BLUE,
    )
    add_panel(
        slide,
        4.95,
        1.95,
        3.9,
        4.8,
        "Nen tang canh bao",
        [
            "Risk warning theo boi canh du an",
            "Evidence + severity + action",
            "Needs human review khi can",
        ],
        TEAL,
    )
    add_panel(
        slide,
        9.1,
        1.95,
        3.4,
        4.8,
        "Nen tang tri thuc",
        [
            "Learning case",
            "Reference snippets",
            "Reusable patterns",
        ],
        GREEN,
    )
    add_footer(slide, "Business capability stack")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "2. Gia tri mo rong lon nhat")
    add_bullets(
        slide,
        [
            "Risk warning theo project context: phat hien som dependency, compliance, release risk.",
            "Learning case hang thang: tong hop best practices va recurring issues de rollout ngang toan cong ty.",
            "Reference snippets: trich nhung doan mo ta tot trong tai lieu input de lam mau chia se.",
            "Universal input: mo duong tu PDF/PPT sang docx, xlsx, URL, wiki, ticket, transcript, va nguon ngoai.",
        ],
        top=2.0,
        height=3.5,
    )
    add_panel(slide, 1.0, 5.7, 5.0, 0.9, "Tac dong", ["Tang chat luong tai lieu, tang chat luong quyet dinh, tang kha nang hoc hoi to chuc"], AMBER)
    add_panel(slide, 6.35, 5.7, 5.0, 0.9, "Luu y", ["Chi duoc dua tren du lieu that; khong invent AI confidence hay risk score gia"], BLUE)
    add_footer(slide, "High-value expansion areas")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "3. Huong mo rong input va nguon du lieu")
    add_bullets(
        slide,
        [
            "Tuong lai khong chi review file upload, ma review normalized content from any source.",
            "Tach document_type khoi input_source_type de khong khoa nghiep vu vao file format.",
            "Ho tro external sources nhu SharePoint, file server, document repository noi bo.",
            "Review chinh thuc tu nguon ngoai nen uu tien snapshot-first de replay va audit duoc.",
        ],
        top=2.0,
        height=3.4,
    )
    add_panel(slide, 1.0, 5.55, 3.5, 1.0, "Input types", ["File, URL, text, JSON, wiki, ticket, transcript"], TEAL)
    add_panel(slide, 4.9, 5.55, 3.2, 1.0, "Connector rule", ["Permission, locator, retrieval log"], BLUE)
    add_panel(slide, 8.45, 5.55, 3.7, 1.0, "Audit rule", ["Snapshot-first cho review chinh thuc"], GREEN)
    add_footer(slide, "Universal input and external source direction")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "4. Thu tu trien khai de khong ngat quang")
    add_panel(slide, 0.8, 2.0, 2.3, 3.8, "P1-P2", ["Governance", "Scope", "Bundle", "Audit nen tang"], BLUE)
    add_panel(slide, 3.35, 2.0, 2.3, 3.8, "P3-P4", ["UI baseline", "Decision", "Risk warning"], TEAL)
    add_panel(slide, 5.9, 2.0, 2.3, 3.8, "P5-P6", ["Learning case", "Snippets", "Pattern sharing"], GREEN)
    add_panel(slide, 8.45, 2.0, 2.3, 3.8, "P7-P8", ["Universal input", "Adapters", "External sources"], AMBER)
    add_panel(slide, 11.0, 2.0, 1.5, 3.8, "Rule", ["Feature flag", "Default path", "Rollback"], NAVY)
    add_bullets(
        slide,
        [
            "Khong expose UI moi khi backend contract chua on dinh.",
            "Moi pha phai co rollback point va minimum safe release gate.",
        ],
        top=6.0,
        height=0.9,
        size=14,
    )
    add_footer(slide, "Low-disruption implementation roadmap")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "5. Dieu kien thanh cong")
    add_bullets(
        slide,
        [
            "Spec, contract, UI, va docs phai thay doi dong bo.",
            "Bundle governance, decision model, va audit contract phai xong truoc khi mo rong learning va multi-source input.",
            "Testing bat buoc cho schema, bundle resolution, risk warning, override, adapter normalization, snapshot/replay, va redaction.",
            "PMO, Backend, Frontend, Design, QA, Security phai co owner ro cho tung pha.",
        ],
        top=2.0,
        height=3.7,
    )
    add_panel(slide, 1.0, 5.8, 5.0, 0.9, "Minimum safe foundation", ["Governance + bundle + runtime contract + UI baseline"], GREEN)
    add_panel(slide, 6.35, 5.8, 5.0, 0.9, "Thong diep chot", ["Mo rong duoc nhung van giu he thong on dinh, audit duoc, va de rollout"], BLUE)
    add_footer(slide, "Success conditions for future business expansion")

    return prs


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs = build()
    prs.save(str(OUT_PATH))
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    main()
