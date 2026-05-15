from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path("artifacts") / "Evaluation_Baseline_Business_Expansion_Safe.pptx"

WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(30, 58, 95)
BLUE = RGBColor(52, 120, 200)
TEAL = RGBColor(39, 144, 132)
GRAY = RGBColor(88, 105, 124)
LIGHT = RGBColor(240, 245, 250)


def set_background(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_title(slide, title, subtitle=""):
    tx = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.8), Inches(1.0))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.runs[0]
        r2.font.name = "Arial"
        r2.font.size = Pt(12)
        r2.font.color.rgb = GRAY


def add_bullets(slide, items, left=0.9, top=1.8, width=11.2, height=4.8, font_size=18):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            run.font.color.rgb = GRAY


def add_two_column(slide, left_title, left_items, right_title, right_items):
    left_box = slide.shapes.add_shape(1, Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.9))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = LIGHT
    left_box.line.color.rgb = BLUE
    right_box = slide.shapes.add_shape(1, Inches(6.75), Inches(1.8), Inches(5.75), Inches(4.9))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = LIGHT
    right_box.line.color.rgb = TEAL

    add_small_heading(slide, left_title, 1.05, 2.05, BLUE)
    add_small_heading(slide, right_title, 7.0, 2.05, TEAL)
    add_bullets(slide, left_items, left=1.0, top=2.5, width=5.0, height=3.9, font_size=14)
    add_bullets(slide, right_items, left=7.0, top=2.5, width=4.9, height=3.9, font_size=14)


def add_small_heading(slide, text, left, top, color):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(4.5), Inches(0.4))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    r = p.runs[0]
    r.font.name = "Arial"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = color


def add_footer(slide, text):
    tx = slide.shapes.add_textbox(Inches(0.7), Inches(7.0), Inches(12.0), Inches(0.25))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    r = p.runs[0]
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(
        slide,
        "Dinh huong mo rong nghiep vu AI Review tai lieu",
        "Ban PowerPoint toi gian de mo on dinh tren PowerPoint desktop va browser viewer.",
    )
    add_bullets(
        slide,
        [
            "Muc tieu: de van hanh, de dung, de audit, de mo rong.",
            "Mo hinh de xuat: (document_type, strictness) -> active evaluation bundle.",
            "Gia tri mo rong: them canh bao rui ro dua tren boi canh thong tin du an.",
        ],
        top=2.0,
        height=2.3,
    )
    add_bullets(
        slide,
        [
            "User chi thao tac voi loai tai lieu, muc danh gia va ket qua review.",
            "Admin quan ly bundle, validate va activate theo scope.",
        ],
        top=4.5,
        height=1.6,
        font_size=16,
    )
    add_footer(slide, "Evaluation Baseline - Summary")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "1. Vi sao can thay doi mo hinh quan ly")
    add_two_column(
        slide,
        "Mo hinh cu",
        [
            "document_type -> evaluation set",
            "Kho mo rong khi mot loai tai lieu co nhieu muc dich review.",
            "Rubric, prompt, policy de bi copy-paste.",
            "Admin kho quan ly bo nao dang active.",
        ],
        "Mo hinh moi",
        [
            "(document_type, strictness) -> active bundle",
            "UI van don gian voi user.",
            "Backend de version, activate, rollback an toan.",
            "Mo duong mo rong objective/context sau nay.",
        ],
    )
    add_footer(slide, "Shift from document-centric configuration to scope-based activation")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "2. Ba muc tieu nghiep vu cua he thong review")
    add_bullets(
        slide,
        [
            "Cai thien chat luong tai lieu: tim thieu sot, diem mo ho, phan can bo sung.",
            "Ho tro quality gate / chot release: dua ra ket qua de PMO va reviewer quyet dinh.",
            "Canh bao rui ro theo boi canh du an: phat hien risk signal dua tren project context va noi dung tai lieu.",
        ],
        top=2.0,
        height=3.0,
    )
    add_bullets(
        slide,
        [
            "He thong khong chi cham diem. He thong phai tao dau ra phuc vu hanh dong tiep theo.",
        ],
        top=5.4,
        height=0.8,
        font_size=16,
    )
    add_footer(slide, "Three operating goals")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "3. Scope va bundle toi gian")
    add_two_column(
        slide,
        "Scope tren UI",
        [
            "Loai tai lieu",
            "Muc danh gia: low / medium / high",
            "Moi scope chi co 1 bundle active tai 1 thoi diem.",
            "Document type mo rong qua master data, strictness mo rong co kiem soat.",
        ],
        "Bundle ben duoi",
        [
            "Rubric version",
            "Prompt version",
            "Policy version",
            "Required rules default",
            "Output schema default",
        ],
    )
    add_footer(slide, "Simple UI, structured backend")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "4. Mo rong Loai tai lieu va Muc danh gia theo kiem soat")
    add_two_column(
        slide,
        "Loai tai lieu",
        [
            "La master data va la truc mo rong nghiep vu chinh.",
            "Moi document_type nen co code, name, description, status.",
            "Cho phep them loai tai lieu moi qua catalog.",
            "Khong encode ngoai le nghiep vu vao ten document_type.",
        ],
        "Muc danh gia",
        [
            "Baseline giu low / medium / high.",
            "Co the doi nhan hien thi thanh Co ban / Tieu chuan / Nghiem ngat.",
            "Chi mo rong khi co nhu cau van hanh ro rang.",
            "Uu tien them metadata truoc khi tao them muc moi.",
        ],
    )
    add_footer(slide, "Controlled expansion principle")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "5. Giam ganh nang cho admin")
    add_two_column(
        slide,
        "Admin can chon",
        [
            "Loai tai lieu",
            "Muc danh gia",
            "Rubric",
            "Prompt",
            "Policy",
            "Ghi chu thay doi",
        ],
        "He thong tu quan ly",
        [
            "Bundle ID",
            "Validation status",
            "Required rules version",
            "Output schema version",
            "Created by / Approved by / Activated by",
        ],
    )
    add_bullets(
        slide,
        ["Luong thay doi chuan: Clone -> Validate -> Activate. Khong sua truc tiep bundle active."],
        top=6.1,
        height=0.6,
        font_size=15,
    )
    add_footer(slide, "Admin simplicity principle")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "6. Vai tro user va admin")
    add_two_column(
        slide,
        "User thuong",
        [
            "Chon Project",
            "Chon Loai tai lieu",
            "Chon Muc danh gia",
            "Upload va review",
            "Xem ket qua va canh bao rui ro",
        ],
        "Admin / PMO",
        [
            "Quan ly scope -> active bundle",
            "Clone bundle hien tai",
            "Sua Rubric / Prompt / Policy khi can",
            "Validate va Activate",
        ],
    )
    add_footer(slide, "Role separation keeps UI easy to learn")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "7. Risk warning dua tren project context")
    add_two_column(
        slide,
        "Project context toi thieu",
        [
            "project_description",
            "project_domain",
            "project_phase",
            "target_milestone",
            "known_constraints",
            "known_dependencies",
            "criticality_level",
        ],
        "Nguyen tac canh bao",
        [
            "Chi canh bao khi co bang chung.",
            "Khong tao risk score gia.",
            "Khong tao AI confidence gia.",
            "Neu thieu du lieu thi phai noi ro khong du thong tin.",
        ],
    )
    add_footer(slide, "Risk warnings must be evidence-based")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "8. Dau ra mong muon cua he thong")
    add_two_column(
        slide,
        "Document assessment",
        [
            "total_score",
            "criteria_scores",
            "summary",
            "strengths",
            "gaps",
            "recommended_improvements",
        ],
        "Risk warnings",
        [
            "title",
            "severity",
            "reason",
            "evidence",
            "recommended_action",
        ],
    )
    add_footer(slide, "Review output should support action and decision")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "9. Learning case va chia se tri thuc")
    add_bullets(
        slide,
        [
            "Khong chi dung review cho tung tai lieu; can tong hop learning case theo thang.",
            "Tong hop best practices, recurring issues, risk patterns, release-readiness cases.",
            "Luong baseline: Review result -> Learning candidates -> Human curation -> Approve -> Publish.",
            "Noi xuat ban: Learning Hub, monthly digest, searchable knowledge base.",
        ],
        top=2.0,
        height=3.2,
    )
    add_bullets(
        slide,
        [
            "AI de xuat, con nguoi duyet truoc khi rollout ngang toan cong ty.",
        ],
        top=5.6,
        height=0.8,
        font_size=16,
    )
    add_footer(slide, "Organizational learning from review result")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "10. Reference snippets tu noi dung tai lieu input")
    add_bullets(
        slide,
        [
            "He thong can hoc ca tu noi dung tai lieu viet tot, khong chi tu ket qua review.",
            "Trich xuat doan requirement, risk, dependency, summary, acceptance criteria viet tot.",
            "Moi snippet can co: why_good, reuse_guidance, document_type, knowledge_area.",
            "Can co masking thong tin nhay cam truoc khi publish thanh snippet library.",
            "Ngoai snippet text, nen co reusable patterns o muc cau truc section.",
        ],
        top=2.0,
        height=3.6,
    )
    add_footer(slide, "Reference snippet and reusable pattern sharing")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "11. Execution roadmap")
    add_bullets(
        slide,
        [
            "P1-P2: governance foundation, scope va bundle foundation.",
            "P3-P4: runtime contract, audit, UI baseline rollout.",
            "P5-P7: risk warning, learning case, reference snippets.",
            "P8-P9: universal input foundation, external source integration.",
        ],
        top=2.0,
        height=3.0,
    )
    add_bullets(
        slide,
        [
            "Rule: chi expose UI moi sau khi backend contract on dinh; moi phase phai co rollback point.",
        ],
        top=5.5,
        height=0.8,
        font_size=16,
    )
    add_footer(slide, "Execution sequence for low-disruption rollout")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "12. Lo trinh trien khai de it rui ro")
    add_bullets(
        slide,
        [
            "Phase 1: scope = (document_type, strictness), global rules/schema, admin quan ly rubric/prompt/policy.",
            "Phase 2: bundle lifecycle day du, compare version, hien thi risk warning block trong UI.",
            "Phase 3: learning case, snippet library, reusable patterns, objective/context_profile neu can.",
        ],
        top=2.0,
        height=3.0,
    )
    add_bullets(
        slide,
        [
            "Nguyen tac chot: UI user phai toi gian; backend co the mo rong dan.",
        ],
        top=5.5,
        height=0.8,
        font_size=16,
    )
    add_footer(slide, "Phased delivery reduces regression risk")

    return prs


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs = build()
    prs.save(str(OUT_PATH))
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    main()
